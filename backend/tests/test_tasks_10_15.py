"""Synthetic electoral acceptance; fixtures do not certify institutional rules."""

import csv
from io import BytesIO, StringIO
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func
from tests.test_auth_rbac import db_session, client, users, headers
from tests.test_tasks_04_09 import scenario, candidate, xlsx, official
from app.models import (
    ElectoralList,
    ListCandidate,
    ElectionRule,
    Candidate,
    Person,
    CandidateValidation,
    AuditLog,
    ListValidation,
    UserModule,
)
from app.services.composition_service import CompositionService
from app.services.affiliate_import_service import AffiliateImportService
from app.services.export_service import render_table, ExportService
from app.services.audit_service import AuditService
from app.utils.enums import ListStatus, ValidationResult


def complete(client, s, local=False):
    row = s["list"]
    if not local:
        rules = {
            **s["rule"]["rules"],
            "positions": [
                {
                    "position": i,
                    "name": f"{'Titular' if i <= 16 else 'Suplente'} {i if i <= 16 else i - 16}",
                    "group": "titular" if i <= 16 else "suplente",
                }
                for i in range(1, 25)
            ],
        }
        r = client.post(
            f"/api/v1/elections/{s['election']['id']}/rules/{s['office']['id']}",
            headers=headers(s["admin"]),
            json=rules,
        )
        assert r.status_code == 201, r.text
        r = client.post(
            f"/api/v1/lists/{row['id']}/rules/adopt", headers=headers(s["admin"])
        )
        assert r.status_code == 200, r.text
        row = r.json()
    if local:
        r = client.post(
            "/api/v1/lists/",
            headers=headers(s["apod"]),
            json={
                **s["list_body"],
                "office_id": s["local"]["id"],
                "municipality_id": s["municipality"]["id"],
                "list_number": "C-01",
            },
        )
        assert r.status_code == 201, r.text
        row = r.json()
    for i in range(22 if local else 24):
        r = client.post(
            f"/api/v1/lists/{row['id']}/candidates",
            headers=headers(s["apod"]),
            json=candidate(
                dni=str(97000000 + (100 if local else 0) + i),
                position=i + 1,
                gender="F" if i % 2 == 0 else "M",
            ),
        )
        assert r.status_code == 201, r.text
    return row


@pytest.mark.parametrize("local", [False, True])
def test_complete_submission_observed_readonly_idempotent(
    client, scenario, db_session, local
):
    s = scenario
    row = complete(client, s, local)
    url = f"/api/v1/lists/{row['id']}"
    h = headers(s["apod"])
    detail = client.get(url, headers=h).json()
    assert (
        detail["composition"]["can_submit"]
        and detail["composition"]["status"] == "pendiente"
    )
    r = client.post(url + "/submit", headers=h).json()
    assert (
        r["submitted"]
        and r["state"] == "enviada_admin"
        and r["approval_blockers"]
        and r["submitted_at"]
    )
    assert client.post(url + "/submit", headers=h).json()["already_submitted"]
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(AuditLog)
            .where(AuditLog.action == "list.submitted")
        )
        == 1
    )
    assert (
        client.put(
            url, headers=h, json={"list_name": "Cambio", "list_number": "Z"}
        ).status_code
        == 409
    )
    c = detail["candidates"][0]
    assert (
        client.put(
            url + f"/candidates/{c['id']}", headers=h, json=candidate()
        ).status_code
        == 409
    )
    assert client.post(url + "/evaluate", headers=h).status_code == 409
    summary = client.get("/api/v1/dashboard/summary", headers=h).json()
    assert (
        summary["sent_lists"] == 1
        and summary["approved_lists"] == 0
        and summary["total_candidates"] == (22 if local else 24)
    )
    review = client.get(
        "/api/v1/admin/candidate-review", headers=headers(s["admin"])
    ).json()
    assert (
        review["total"] == len(detail["candidates"])
        and review["items"][0]["lists"][0]["id"] == row["id"]
    )


def test_incomplete_evaluation_persists_and_can_correct(client, scenario, db_session):
    s = scenario
    url = f"/api/v1/lists/{s['list']['id']}"
    h = headers(s["apod"])
    r = client.post(url + "/submit", headers=h).json()
    assert not r["submitted"] and r["state"] == "incompleta"
    assert db_session.scalar(select(func.count()).select_from(ListValidation)) == 1
    assert not db_session.scalar(
        select(AuditLog).where(AuditLog.action == "list.submitted")
    )
    assert (
        client.post(url + "/candidates", headers=h, json=candidate()).status_code == 201
    )
    assert client.get(url, headers=h).json()["status"] == "borrador"


@pytest.mark.parametrize(
    "defect,code",
    [
        ("missing", "COUNT"),
        ("extra", "EXTRA"),
        ("role", "POSITION_ROLE"),
        ("gender", "PARITY"),
        ("unknown", "GENDER_POLICY"),
        ("context", "CONTEXT"),
        ("duplicate", "DUPLICATE_DNI"),
        ("position", "DUPLICATE_POSITION"),
    ],
)
def test_composition_defects(client, scenario, db_session, defect, code):
    s = scenario
    row = complete(client, s)
    obj = db_session.get(ElectoralList, row["id"])
    members = list(
        db_session.scalars(
            select(ListCandidate)
            .where(ListCandidate.list_id == obj.id)
            .order_by(ListCandidate.position_number)
        )
    )
    first = members[0]
    c = db_session.get(Candidate, first.candidate_id)
    p = db_session.get(Person, c.person_id)
    if defect == "missing":
        db_session.delete(first)
    elif defect == "extra":
        first.position_number = 25
    elif defect == "role":
        first.cargo_group = "suplente-inválido"
    elif defect in ("gender", "unknown"):
        first.gender = p.gender = "M" if defect == "gender" else "X"
    elif defect == "context":
        c.office_id = s["local"]["id"]
    elif defect == "duplicate":
        members[1].candidate_id = first.candidate_id
    elif defect == "position":
        members[1].position_number = first.position_number
    db_session.flush()
    result = CompositionService(db_session).evaluate(obj)
    assert not result["can_submit"] and code in {i["code"] for i in result["issues"]}


def test_alternation_applies_to_council_not_deputies(client, scenario, db_session):
    s = scenario
    for local in [False, True]:
        row = complete(client, s, local)
        members = list(
            db_session.scalars(
                select(ListCandidate)
                .where(ListCandidate.list_id == row["id"])
                .order_by(ListCandidate.position_number)
            )
        )
        for m in members[:2]:
            p = db_session.get(
                Person, db_session.get(Candidate, m.candidate_id).person_id
            )
            p.gender = m.gender = "M" if m.position_number == 1 else "F"
        db_session.flush()
        result = CompositionService(db_session).evaluate(
            db_session.get(ElectoralList, row["id"])
        )
        assert result["can_submit"] is (not local)
        assert ("ALTERNATION" in {i["code"] for i in result["issues"]}) is local


def test_approval_without_identity_provider_uses_current_register(
    client, scenario, db_session
):
    s = scenario
    row = complete(client, s)
    rule = db_session.get(ElectionRule, row["rule_version_id"])
    rule.rules = {
        **rule.rules,
        "template_is_test": False,
        "other_requirements_confirmed": True,
        "requires_renaper": True,  # Historical configuration must have no effect.
    }
    db_session.commit()
    AffiliateImportService(db_session).import_excel(
        xlsx([official(**{"Matrícula": str(97000000 + i)}) for i in range(24)]),
        "synthetic-register.xlsx",
        s["admin"].id,
    )
    r = client.post(f"/api/v1/lists/{row['id']}/submit", headers=headers(s["apod"]))
    assert r.status_code == 200, r.text
    assert (
        r.json()["state"] == "aprobada_sistema" and r.json()["approval_blockers"] == []
    )
    events = list(
        db_session.scalars(
            select(AuditLog)
            .where(
                AuditLog.action.in_(["list.submitted", "list.automatically_approved"])
            )
            .order_by(AuditLog.id)
        )
    )
    assert (
        len(events) == 2
        and events[0].user_id == s["apod"].id
        and events[1].user_id is None
    )
    assert events[1].details_json["submission_event_id"] == events[0].id
    v = db_session.scalar(select(ListValidation))
    assert (
        v.response_json["approval_evaluated"]
        and v.response_json["resulting_state"] == "aprobada_sistema"
    )


def test_send_rechecks_old_ok_and_rolls_back_failed_audit(
    client, scenario, db_session, monkeypatch
):
    s = scenario
    row = complete(client, s)
    for v in db_session.scalars(select(CandidateValidation)):
        v.status = ValidationResult.OK
        v.response_json = {"origin": "authorized_provider", "approvable": True}
    db_session.commit()
    original = AuditService.record

    def fail(self, actor, action, *args, **kwargs):
        if action == "list.submitted":
            raise RuntimeError("synthetic audit failure")
        return original(self, actor, action, *args, **kwargs)

    monkeypatch.setattr(AuditService, "record", fail)
    with pytest.raises(RuntimeError):
        client.post(f"/api/v1/lists/{row['id']}/submit", headers=headers(s["apod"]))
    assert db_session.get(ElectoralList, row["id"]).status == ListStatus.BORRADOR
    assert db_session.scalar(select(func.count()).select_from(ListValidation)) == 0
    monkeypatch.setattr(AuditService, "record", original)
    result = client.post(
        f"/api/v1/lists/{row['id']}/submit", headers=headers(s["apod"])
    ).json()
    assert result["state"] == "enviada_admin" and result["approval_blockers"]
    assert any(
        v.status == ValidationResult.PENDIENTE
        for v in db_session.scalars(select(CandidateValidation))
    )


def test_reports_full_filtered_scope_and_exports(client, scenario, db_session):
    s = scenario
    h = headers(s["admin"])
    for i in range(30):
        r = client.post(
            "/api/v1/lists/",
            headers=h,
            json={
                **s["list_body"],
                "list_name": f"Filtro Ágil {i}",
                "list_number": f"R-{i}",
            },
        )
        assert r.status_code == 201, r.text
    qs = f"search=Filtro&office_id={s['office']['id']}&election_id={s['election']['id']}&status=borrador"
    data = client.get("/api/v1/admin/lists?" + qs, headers=h).json()
    assert data["total"] == 30 and len(data["items"]) == 25
    assert (
        len(
            client.get("/api/v1/admin/lists?" + qs + "&page=2", headers=h).json()[
                "items"
            ]
        )
        == 5
    )
    summary = client.get("/api/v1/admin/reports?" + qs, headers=h).json()
    assert (
        summary["total_lists"] == 30
        and summary["approval_rate"] == 0
        and summary["by_office"][0]["lists"] == 30
    )
    for fmt in ["xlsx", "csv", "pdf"]:
        kind = "reports" if fmt == "pdf" else "lists"
        r = client.get(f"/api/v1/admin/exports/{kind}?format={fmt}&{qs}", headers=h)
        assert r.status_code == 200, r.text
        assert f".{fmt}" in r.headers["content-disposition"]
        if fmt == "csv":
            assert len(list(csv.reader(StringIO(r.content.decode("utf-8-sig"))))) == 31
        elif fmt == "xlsx":
            assert len(list(load_workbook(BytesIO(r.content))["Datos"].values)) == 31
        else:
            assert r.content.startswith(b"%PDF")
    assert (
        db_session.scalar(
            select(func.count())
            .select_from(AuditLog)
            .where(AuditLog.action == "export.generated")
        )
        == 3
    )
    empty = client.get("/api/v1/admin/reports?search=NO-MATCH", headers=h).json()
    assert empty["total_lists"] == 0 and empty["approval_rate"] == 0
    assert (
        client.get(
            "/api/v1/admin/exports/reports?format=pdf&search=NO-MATCH", headers=h
        ).status_code
        == 200
    )


def test_export_text_and_formula_safety():
    columns = [("dni", "DNI"), ("name", "Nombre")]
    rows = [
        {"dni": "00123456", "name": '=HYPERLINK("bad")'},
        {"dni": "0099", "name": 'Álvarez, "Prueba"\nsegunda línea'},
        {"dni": "0042", "name": " \t@SUM(1)"},
    ]
    book = load_workbook(BytesIO(render_table(rows, columns, "xlsx", "Prueba")))
    assert (
        book["Datos"]["A2"].value == "00123456" and book["Datos"]["B2"].data_type == "s"
    )
    values = list(
        csv.reader(
            StringIO(render_table(rows, columns, "csv", "Prueba").decode("utf-8-sig"))
        )
    )
    assert (
        values[1][0] == "00123456"
        and values[1][1].startswith("'=")
        and values[3][1].startswith("'")
    )
    assert values[2][1] == rows[1]["name"]


@pytest.mark.parametrize(
    "path",
    [
        "/admin/lists",
        "/admin/reports",
        "/admin/candidate-review",
        "/admin/audit",
        "/admin/audit/1",
        "/admin/exports/lists",
        "/admin/exports/reports?format=pdf",
        "/admin/exports/padron",
    ],
)
def test_admin_only(client, users, path):
    assert client.get("/api/v1" + path).status_code == 401
    assert client.get("/api/v1" + path, headers=headers(users[1])).status_code == 403


def test_dashboard_scope_revocation_and_all_states(client, scenario, db_session):
    s = scenario
    h = headers(s["admin"])
    a = headers(s["apod"])
    for i, status in enumerate(ListStatus):
        row = ElectoralList(
            **{**s["list_body"], "list_number": f"State-{i}"},
            created_by=s["admin"].id,
            status=status,
        )
        db_session.add(row)
    db_session.commit()
    global_data = client.get("/api/v1/dashboard/summary", headers=h).json()
    assert (
        global_data["total_lists"] == 7
        and global_data["approved_lists"] == 1
        and global_data["approval_denominator"] == 7
    )
    assert (
        global_data["approval_rate"] == 14.29 and global_data["active_apoderados"] == 1
    )
    own = client.get("/api/v1/dashboard/summary", headers=a).json()
    assert (
        own["total_lists"] == 1
        and own["active_apoderados"] is None
        and len(own["modules"]) == 2
    )
    for m in db_session.scalars(
        select(UserModule).where(UserModule.user_id == s["apod"].id)
    ):
        m.enabled = False
    db_session.commit()
    revoked = client.get("/api/v1/dashboard/summary", headers=a).json()
    assert revoked["total_lists"] == 0 and revoked["modules"] == []
    assert (
        client.post(f"/api/v1/lists/{s['list']['id']}/submit", headers=a).status_code
        == 403
    )


def test_audit_filters_date_context_safe_detail(client, scenario, db_session):
    s = scenario
    h = headers(s["admin"])
    lid = s["list"]["id"]
    r = client.post(
        f"/api/v1/lists/{lid}/candidates", headers=headers(s["apod"]), json=candidate()
    )
    assert r.status_code == 201
    event = AuditService(db_session).record(
        s["admin"].id,
        "safe.test",
        "electoral_lists",
        lid,
        {
            "password": "never-store",
            "token": "never-store",
            "before": {"name": "Prueba", "database_url": "never-store"},
            "url": "postgresql://user:pass@private/db",
        },
    )
    db_session.commit()
    assert "never-store" not in str(event.details_json) and "user:pass" not in str(
        event.details_json
    )
    data = client.get(
        f"/api/v1/admin/audit?list_id={lid}&page_size=2", headers=h
    ).json()
    assert (
        data["total"] >= 3
        and len(data["items"]) == 2
        and data["items"][0]["id"] > data["items"][1]["id"]
    )
    selected = client.get(
        f"/api/v1/admin/audit?actor_id={s['admin'].id}&action=safe.test&entity_type=electoral_lists&entity_id={lid}",
        headers=h,
    ).json()
    assert (
        selected["total"] == 1
        and selected["items"][0]["actor_name"] == s["admin"].full_name
    )
    assert (
        client.get(
            "/api/v1/admin/audit?date_from=2026-10-01&date_to=2026-01-01", headers=h
        ).status_code
        == 422
    )
    assert (
        client.get(f"/api/v1/admin/audit/{event.id}", headers=h).json()["details"]
        == event.details_json
    )
    assert (
        client.delete(f"/api/v1/admin/audit/{event.id}", headers=h).status_code == 405
    )


def test_export_limit_explicit_not_truncated(client, scenario, monkeypatch):
    monkeypatch.setattr(ExportService, "maximum", 0)
    assert (
        client.get(
            "/api/v1/admin/exports/lists", headers=headers(scenario["admin"])
        ).status_code
        == 413
    )


def test_padron_export_filtered_fields_and_safe_text(client, scenario):
    import pandas as pd

    s = scenario
    h = headers(s["admin"])
    buf = BytesIO()
    pd.DataFrame(
        [
            {
                "Matrícula": "00123456",
                "Nombre": "Prueba",
                "Apellido": "Álvarez",
                "Sección": "Uno",
                "Circuito": "A",
                "Domicilio": "=DANGER()",
            },
            {
                "Matrícula": "00123457",
                "Nombre": "Otra",
                "Apellido": "Sintética",
                "Sección": "Dos",
                "Circuito": "B",
            },
        ]
    ).to_excel(buf, index=False)
    book = load_workbook(BytesIO(buf.getvalue()))
    book.active.cell(2, 6).value = "=DANGER()"
    book.active.cell(2, 6).data_type = "s"
    buf = BytesIO()
    book.save(buf)
    r = client.post(
        "/api/v1/padron/import",
        headers=h,
        files={
            "file": (
                "synthetic.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 200, r.text
    visible = client.get("/api/v1/padron?section=Uno", headers=h).json()
    assert visible["total"] == 1
    assert visible["items"][0]["last_name"] == "Álvarez"
    for fmt in ["csv", "xlsx"]:
        r = client.get(
            f"/api/v1/admin/exports/padron?format={fmt}&section=Uno", headers=h
        )
        assert r.status_code == 200, r.text
        if fmt == "xlsx":
            sheet = load_workbook(BytesIO(r.content))["Datos"]
            values = list(sheet.values)
            assert sheet.cell(2, 18).data_type == "s"
        else:
            values = list(csv.reader(StringIO(r.content.decode("utf-8-sig"))))
        assert len(values) == 2 and len(values[0]) == 18 and values[1][8] == "00123456"


def test_pdf_extremely_long_assignment_wraps_pages():
    from app.services.export_service import LIST_COLUMNS

    content = render_table(
        [
            {
                "list_name": "Sintética",
                "apoderados_text": "; ".join(
                    "Apoderado ficticio " + str(i) for i in range(250)
                ),
            }
        ],
        LIST_COLUMNS,
        "pdf",
        "Prueba",
    )
    assert content.startswith(b"%PDF")


def test_review_ignores_superseded_legacy_warning(client, scenario, db_session):
    from app.utils.enums import ValidationType

    s = scenario
    h = headers(s["admin"])
    r = client.post(
        f"/api/v1/lists/{s['list']['id']}/candidates",
        headers=headers(s["apod"]),
        json=candidate(),
    )
    cid = r.json()["id"]
    db_session.add(
        CandidateValidation(
            candidate_id=cid,
            candidate_revision=None,
            validation_type=ValidationType.AFILIACION,
            status=ValidationResult.WARNING,
            message="Legacy",
        )
    )
    for v in db_session.scalars(
        select(CandidateValidation).where(CandidateValidation.candidate_id == cid)
    ):
        if v.candidate_revision is not None:
            v.status = ValidationResult.OK
    db_session.commit()
    assert client.get("/api/v1/admin/candidate-review", headers=h).json()["total"] == 0
    assert not client.get("/api/v1/candidates/review", headers=h).json()


@pytest.mark.parametrize(
    "local,delta", [(False, -1), (False, 1), (True, -1), (True, 1)]
)
def test_exact_cardinality_both_offices(client, scenario, db_session, local, delta):
    row = complete(client, scenario, local)
    obj = db_session.get(ElectoralList, row["id"])
    first = db_session.scalar(
        select(ListCandidate).where(ListCandidate.list_id == obj.id)
    )
    if delta < 0:
        db_session.delete(first)
    else:
        db_session.add(
            ListCandidate(
                list_id=obj.id,
                candidate_id=first.candidate_id,
                position_number=23 if local else 25,
                gender=first.gender,
                cargo_group=first.cargo_group,
                cargo_label="Extra",
            )
        )
    db_session.flush()
    result = CompositionService(db_session).evaluate(obj)
    assert result["details"]["current_positions"] == (22 if local else 24) + delta
    assert "COUNT" in {i["code"] for i in result["issues"]} and not result["can_submit"]


def test_submission_outside_window_and_other_assignee(client, scenario, db_session):
    from datetime import date, timedelta
    from app.models import User, Election
    from app.utils.enums import UserRole, UserModuleType

    s = scenario
    other = User(
        username="other-synthetic",
        email="other-synthetic@example.com",
        full_name="Otra cuenta",
        role=UserRole.APODERADO,
        is_active=True,
        password_hash=s["apod"].password_hash,
    )
    db_session.add(other)
    db_session.flush()
    db_session.add(
        UserModule(
            user_id=other.id,
            module_type=UserModuleType.DIPUTADOS_PROVINCIALES,
            election_id=s["election"]["id"],
            office_id=s["office"]["id"],
            enabled=True,
        )
    )
    db_session.commit()
    url = f"/api/v1/lists/{s['list']['id']}/submit"
    assert client.post(url, headers=headers(other)).status_code == 403
    election = db_session.get(Election, s["election"]["id"])
    election.loading_closes = date.today() - timedelta(days=1)
    db_session.commit()
    assert client.post(url, headers=headers(s["apod"])).status_code == 409
    assert not db_session.scalar(
        select(AuditLog).where(AuditLog.action == "list.submitted")
    )


def test_audit_dates_use_argentina_calendar(client, scenario, db_session):
    from datetime import datetime

    s = scenario
    for time in (datetime(2026, 1, 2, 1), datetime(2026, 1, 2, 4)):
        db_session.add(
            AuditLog(
                user_id=s["admin"].id,
                action="date.test",
                entity_type="electoral_lists",
                entity_id=s["list"]["id"],
                created_at=time,
            )
        )
    db_session.commit()
    result = client.get(
        "/api/v1/admin/audit?action=date.test&date_from=2026-01-01&date_to=2026-01-01",
        headers=headers(s["admin"]),
    ).json()
    assert result["total"] == 1 and result["items"][0]["created_at"].startswith(
        "2026-01-02T01"
    )


def test_legacy_candidate_creation_and_audit_are_atomic(
    client, scenario, db_session, monkeypatch
):
    s = scenario
    body = {
        "person": {
            k: v
            for k, v in candidate(dni="96123456").items()
            if k not in ("position", "action")
        },
        "office_id": s["office"]["id"],
        "election_id": s["election"]["id"],
    }
    previous = db_session.scalar(select(func.count()).select_from(Person))
    original = AuditService.record

    def fail(*args, **kwargs):
        raise RuntimeError("Synthetic audit failure")

    monkeypatch.setattr(AuditService, "record", fail)
    with pytest.raises(RuntimeError):
        client.post("/api/v1/candidates", headers=headers(s["admin"]), json=body)
    assert db_session.scalar(select(func.count()).select_from(Person)) == previous
    assert not db_session.scalar(select(Candidate))
    monkeypatch.setattr(AuditService, "record", original)
    result = client.post("/api/v1/candidates", headers=headers(s["admin"]), json=body)
    assert result.status_code == 201, result.text
    event = db_session.scalar(
        select(AuditLog).where(
            AuditLog.entity_type == "candidates", AuditLog.action == "candidate.created"
        )
    )
    assert (
        event.user_id == s["admin"].id
        and event.details_json["affiliation_status"] == "warning"
    )
