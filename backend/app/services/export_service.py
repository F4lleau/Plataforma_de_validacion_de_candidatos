from app.models import Election, Office, Municipality, User
import csv
from io import BytesIO, StringIO
from datetime import datetime, timezone
from xml.sax.saxutils import escape
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from app.services.reporting_service import ReportingService
from app.services.padron_service import PadronService

STATE_LABELS = {
    "borrador": "Borrador",
    "incompleta": "Incompleta",
    "en_validacion": "En validación",
    "rechazada_composicion": "Rechazada por composición",
    "enviada_admin": "Enviada al Admin",
    "aprobada_sistema": "Aprobada por Sistema",
}

LIST_COLUMNS = [
    ("list_number", "Número"),
    ("list_name", "Lista"),
    ("election_name", "Elección"),
    ("office_name", "Cargo"),
    ("municipality_name", "Localidad"),
    ("candidate_count", "Candidatos"),
    ("status", "Estado"),
    ("apoderados_text", "Apoderados"),
]
PADRON_COLUMNS = [
    ("section", "Sección"),
    ("section_code", "Cod. Sección"),
    ("circuit", "Circuito"),
    ("circuit_code", "Cod. Circuito"),
    ("last_name", "Apellido"),
    ("first_name", "Nombre"),
    ("gender", "Género"),
    ("document_type", "Tipo documento"),
    ("dni", "Matrícula"),
    ("birth_date", "Fecha nacimiento"),
    ("birth_class", "Clase"),
    ("elector_status", "Estado elector"),
    ("affiliation_status", "Estado afiliación"),
    ("affiliation_date", "Fecha afiliación"),
    ("illiterate", "Analfabeto"),
    ("profession", "Profesión"),
    ("address_date", "Fecha domicilio"),
    ("address", "Domicilio"),
]
MIME = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def text_value(value):
    if value is None:
        return ""
    return str(getattr(value, "value", value))


def csv_safe(value):
    text = text_value(value)
    # Spreadsheet applications may ignore leading whitespace before a formula.
    return (
        "'" + text
        if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r"))
        else text
    )


def render_table(rows, columns, fmt, title, subtitle="", metrics=None):
    if fmt == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow([label for _, label in columns])
        for row in rows:
            writer.writerow([csv_safe(row.get(key)) for key, _ in columns])
        return output.getvalue().encode("utf-8-sig")
    if fmt == "xlsx":
        book = Workbook(write_only=True)
        sheet = book.create_sheet("Datos")
        sheet.append([label for _, label in columns])
        for row in rows:
            values = []
            for key, _ in columns:
                value = row.get(key)
                if isinstance(value, (int, float)):
                    values.append(value)
                else:
                    cell = WriteOnlyCell(sheet, value=text_value(value))
                    cell.data_type = "s"
                    values.append(cell)
            sheet.append(values)
        meta = book.create_sheet("Información")
        meta.append(["Reporte", title])
        meta.append(["Filtros", subtitle])
        meta.append(["Registros", len(rows)])
        meta.append(["Generado UTC", datetime.now(timezone.utc).isoformat()])
        if metrics:
            for key, label in [
                ("total_lists", "Total listas"),
                ("total_candidates", "Candidatos"),
                ("active_apoderados", "Apoderados activos"),
                ("approval_rate", "Aprobación (%)"),
                ("approval_denominator", "Denominador aprobación"),
            ]:
                meta.append([label, metrics.get(key)])
        output = BytesIO()
        book.save(output)
        return output.getvalue()
    if fmt != "pdf":
        raise HTTPException(422, "Formato no disponible.")
    output = BytesIO()
    page = landscape(A4)
    styles = getSampleStyleSheet()
    small = ParagraphStyle(
        "Cell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=7,
        leading=9,
        wordWrap="CJK",
    )
    flow = [
        Paragraph(escape(title), styles["Title"]),
        Paragraph(escape(subtitle or "Sin filtros"), styles["Normal"]),
        Spacer(1, 10),
    ]
    if metrics:
        flow.extend(
            [
                Paragraph(
                    f"Listas: {metrics['total_lists']} | Candidatos: {metrics['total_candidates']} | Aprobación: {metrics['approval_rate']}% ({metrics['approved_lists']}/{metrics['approval_denominator']} listas)",
                    styles["Normal"],
                ),
                Spacer(1, 8),
            ]
        )
    flow.extend(
        [
            Paragraph(
                f"Registros: {len(rows)}. Los datos de prueba y RENAPER pendiente no equivalen a aprobación institucional.",
                styles["Normal"],
            ),
            Spacer(1, 12),
        ]
    )
    data = [[Paragraph(escape(label), small) for _, label in columns]]
    for row in rows:
        data.append(
            [
                Paragraph(
                    escape(text_value(row.get(key))).replace("\n", "<br/>"), small
                )
                for key, _ in columns
            ]
        )
    if len(data) == 1:
        flow.append(Paragraph("Sin resultados para estos filtros.", styles["Normal"]))
    else:
        widths = [45, 120, 100, 95, 90, 60, 90, 100] if len(columns) == 8 else None
        table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT", splitInRow=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE9F6")),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F5F7FA")],
                    ),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#224F82")),
                ]
            )
        )
        flow.append(table)

    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(
            28,
            18,
            "Junta Electoral - Reporte local | Generado "
            + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        )
        canvas.drawRightString(page[0] - 28, 18, f"Página {doc.page}")

    SimpleDocTemplate(
        output,
        pagesize=page,
        leftMargin=28,
        rightMargin=28,
        topMargin=26,
        bottomMargin=32,
        title=title,
        author="Junta Electoral",
    ).build(flow, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


class ExportService(ReportingService):
    maximum = 50000

    def lists(self, user, filters, fmt, report=False):
        self.check_filters(user, filters)
        # Obtain full filtered scope once, never silently export just the current UI page.
        items, total = self.reporting.list_page(user, filters, 1, self.maximum + 1)
        if total > self.maximum:
            raise HTTPException(
                413, "El resultado supera 50000 registros. Acote los filtros."
            )
        if fmt == "pdf" and total > 5000:
            raise HTTPException(
                413, "PDF admite hasta 5000 listas; acote los filtros o use Excel/CSV."
            )
        for item in items:
            item["apoderados_text"] = "; ".join(a["name"] for a in item["apoderados"])
        for item in items:
            item["status"] = STATE_LABELS.get(
                text_value(item["status"]), text_value(item["status"])
            )
        scope = filters.model_dump(mode="json", exclude_none=True)
        labels = []
        for key, model, label in [
            ("election_id", Election, "Elección"),
            ("office_id", Office, "Cargo"),
            ("municipality_id", Municipality, "Localidad"),
            ("apoderado_id", User, "Apoderado"),
        ]:
            if scope.get(key):
                resource = self.require(model, scope[key])
                labels.append(
                    f"{label}: {resource.full_name if model == User else resource.name}"
                )
        if scope.get("status"):
            labels.append("Estado: " + STATE_LABELS[scope["status"]])
        if scope.get("search"):
            labels.append("Búsqueda: " + scope["search"])
        metrics = self.summary(user, filters)
        content = render_table(
            items,
            LIST_COLUMNS,
            fmt,
            "Reporte electoral" if report else "Listas electorales",
            ", ".join(labels),
            metrics,
        )
        self.audit.record(
            user.id,
            "export.generated",
            "electoral_lists",
            None,
            {
                "format": fmt,
                "kind": "report" if report else "lists",
                "filters": scope,
                "rows": total,
            },
        )
        self.commit()
        return (
            content,
            MIME[fmt],
            f"{'reporte' if report else 'listas'}-electorales.{fmt}",
        )

    def padron(self, user, filters, fmt):
        result = PadronService(self.db).query(
            **filters, page=1, page_size=self.maximum + 1
        )
        if result["total"] > self.maximum:
            raise HTTPException(
                413, "El resultado supera 50000 registros. Acote los filtros."
            )
        rows = [
            {key: getattr(m, key) for key, _ in PADRON_COLUMNS} for m in result["items"]
        ]
        content = render_table(
            rows,
            PADRON_COLUMNS,
            fmt,
            "Padrón de afiliados",
            ", ".join(f"{k}: {v}" for k, v in filters.items() if v),
        )
        self.audit.record(
            user.id,
            "export.generated",
            "party_members",
            None,
            {
                "format": fmt,
                "kind": "padron",
                "rows": result["total"],
                "batch_id": result["current_batch"].id
                if result["current_batch"]
                else None,
            },
        )
        self.commit()
        return content, MIME[fmt], f"padron-afiliados.{fmt}"
